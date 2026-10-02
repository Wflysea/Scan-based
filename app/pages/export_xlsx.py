"""导出 xlsx 备份：将库存与出入库流水导出为 Excel（openpyxl + FilePicker）。

字段范围：
- Sheet「商品库存」：编码, 商品名称, 当前库存, 安全库存, 单位, 仓库/库位, 更新时间
- Sheet「出入库流水」：序号, 编码, 商品名称, 操作类型, 数量, 操作人, 仓库/库位, 备注, 操作时间

文件命名：库存备份_YYYYMMDD_HHMMSS.xlsx（例如 库存备份_20261002_205000.xlsx）

交互逻辑：点击按钮 → 弹出系统保存对话框（已预填文件名）→ 用户选择路径或取消；
  选择后写入并弹成功提示（含完整路径）；取消则不导出；异常弹错误提示。
"""
import datetime as _dt

import flet as ft

from app import db
from app.config import PRIMARY, OUT_COLOR, BTN_HEIGHT


def build_export_button(app):
    page = app.page
    # FilePicker 在页面 overlay 上只需挂载一次（单例），多次进入首页复用同一实例
    fp = getattr(page, "_export_fp", None)
    if fp is None:
        fp = ft.FilePicker(on_result=lambda e: _on_picked(app, e))
        page.overlay.append(fp)
        page._export_fp = fp

    return ft.FilledTonalButton(
        "导出 xlsx 备份", icon=ft.Icons.DOWNLOAD, height=BTN_HEIGHT,
        on_click=lambda e: _start_export(fp),
    )


def _start_export(file_picker: ft.FilePicker):
    name = f"库存备份_{_dt.datetime.now():%Y%m%d_%H%M%S}.xlsx"
    file_picker.save_file(dialog_title="导出库存备份", file_name=name)


def _on_picked(app, e: ft.FilePickerResultEvent):
    if not e.path:
        return  # 用户取消
    app.snack("正在导出备份…", PRIMARY)
    try:
        _write_xlsx(e.path)
        app.snack(f"已导出备份：{e.path}", PRIMARY)
    except Exception as ex:  # noqa: BLE001
        app.snack(f"导出失败：{ex}", OUT_COLOR)


def _write_xlsx(path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()

    # ---- Sheet 1：商品库存 ----
    ws1 = wb.active
    ws1.title = "商品库存"
    ws1.append(["编码", "商品名称", "当前库存", "安全库存", "单位", "仓库/库位", "更新时间"])
    for p in db.list_products():
        ws1.append([
            p["code"], p["name"], p["quantity"], p["safety_stock"],
            p["unit"], p["location"], p["updated_at"],
        ])

    # ---- Sheet 2：出入库流水 ----
    ws2 = wb.create_sheet("出入库流水")
    ws2.append(["序号", "编码", "商品名称", "操作类型", "数量", "操作人", "仓库/库位", "备注", "操作时间"])
    for t in db.list_transactions():
        ws2.append([
            t["id"], t["code"], t["name"],
            "入库" if t["op_type"] == "IN" else "出库",
            t["quantity"], t["operator"], t["location"], t["note"], t["created_at"],
        ])

    # ---- 表头样式 + 列宽 ----
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1565C0")
    for ws in (ws1, ws2):
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = 18
        ws.freeze_panes = "A2"

    wb.save(path)
